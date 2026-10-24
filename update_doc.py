# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

SRC = r'C:\Users\ASUS\Downloads\《企业级应用开发 》课程大作业任务书.docx'
DST = r'C:\Users\ASUS\Downloads\《企业级应用开发 》课程大作业任务书_updated.docx'

doc = Document(SRC)

LQ = '“'  # "
RQ = '”'  # "


def make_paragraph(text=''):
    if text:
        xml = '<w:p {ns}><w:r><w:rPr><w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'.format(
            ns=nsdecls('w'), text=text)
    else:
        xml = '<w:p {ns}><w:r><w:rPr><w:sz w:val="21"/><w:szCs w:val="21"/></w:rPr><w:t></w:t></w:r></w:p>'.format(
            ns=nsdecls('w'))
    return parse_xml(xml)


def make_bold_paragraph(text):
    xml = '''<w:p {ns}>
        <w:r>
            <w:rPr>
                <w:b/>
                <w:sz w:val="21"/>
                <w:szCs w:val="21"/>
            </w:rPr>
            <w:t xml:space="preserve">{text}</w:t>
        </w:r>
    </w:p>'''.format(ns=nsdecls('w'), text=text)
    return parse_xml(xml)


def insert_elements_after(after_elem, elements):
    current = after_elem
    for elem in elements:
        current.addnext(elem)
        current = elem


def add_img_to_para(image_path, width_inches=5.5):
    p = doc.add_paragraph()
    r = p.add_run()
    r.add_picture(image_path, width=Inches(width_inches))
    return p._element


# ============================================================
# Working bottom-up: 267, 257, 240, 170, 152, 73
# ============================================================

para_267 = doc.paragraphs[267]._element
para_257 = doc.paragraphs[257]._element
para_240 = doc.paragraphs[240]._element
para_170 = doc.paragraphs[170]._element
para_152 = doc.paragraphs[152]._element
para_73  = doc.paragraphs[73]._element

# ============================================================
# 6.3 & 6.4
# ============================================================

exp_comm_title = '6.3 社区互动模块心得体会'
exp_comm_1 = (
    '社区互动模块是项目中业务逻辑最密集、关联最多的模块之一。'
    '实现 Feed 流推荐时，推模式（发帖时直接写入所有粉丝收件箱）在大V 场景下存在明显的写放大问题——'
    '拥有百万粉丝的用户发一条动态就需要写百万次数据库。'
    '通过引入 BigVConfig 配置表，粉丝数超过阈值的用户采用拉模式（粉丝主动拉取），'
    '普通用户继续使用推模式，推拉结合的分发策略在系统吞吐与数据实时性之间取得了良好平衡。'
    '这让我深刻理解了经典系统设计问题在真实业务中的取舍过程。'
)
exp_comm_2 = (
    'Elasticsearch 全文搜索的集成也积累了大量实践经验。'
    '从索引结构设计（字段类型选择、analyzer 配置）、数据同步策略（事件驱动 + 定时补偿）、'
    '到查询语法优化（bool 查询、should/filter 上下文、高亮显示），'
    '完整走通了一套搜索引擎从 0 到 1 的落地流程。'
    '同时也意识到 ES 并非银弹——不合理的索引设计会导致查询性能下降，'
    '因此为不同业务场景设计了差异化的索引映射和查询模板。'
)

exp_msg_title = '6.4 消息与私信模块心得体会'
exp_msg_1 = (
    '消息与私信模块让我深入理解了 WebSocket 长连接管理与消息推送的工程实践。'
    '在实现实时私信功能时，最开始采用短轮询方案，每隔几秒查询一次新消息，'
    '不仅浪费服务器资源，消息延迟也较高。'
    '后期升级为 WebSocket 长连接，通过统一的心跳检测、断线重连机制维持连接稳定性。'
    '但随之而来的是连接状态管理问题——如何准确追踪每个用户的在线状态、如何处理连接断开后的消息补发，'
    '这些都是实际生产中必须面对的挑战。'
)
exp_msg_2 = (
    'RocketMQ 的引入也让我对消息队列在系统解耦中的作用有了真实体感。'
    '此前消息推送、数据统计等后置任务全部同步执行，接口响应耗时集中在 200-500ms。'
    '将这些非核心链路剥离到 RocketMQ 异步消费后，接口耗时降至 20ms 以内，系统吞吐量显著提升。'
    '同时，消息可靠性（生产者重试、消费者重试、死信队列）和顺序消费的配置也让我积累了'
    '消息中间件在生产环境中的落地经验。'
)

elements_6 = [
    make_paragraph(),
    make_bold_paragraph(exp_comm_title),
    make_paragraph(),
    make_paragraph(exp_comm_1),
    make_paragraph(),
    make_paragraph(exp_comm_2),
    make_paragraph(),
    make_bold_paragraph(exp_msg_title),
    make_paragraph(),
    make_paragraph(exp_msg_1),
    make_paragraph(),
    make_paragraph(exp_msg_2),
]
insert_elements_after(para_267, elements_6)
print('6.3 & 6.4 inserted!')

# ============================================================
# 5.3 & 5.4
# ============================================================

conclusion_comm_title = '5.3 社区互动模块结论与展望'
conclusion_comm_text_title = '结论'
conclusion_comm_content = (
    '社区互动模块聚焦用户动态发布、Feed 流推荐、点赞评论、话题运营和全局搜索等社交核心场景，'
    '实现了从内容生产、分发到互动反馈的完整闭环。'
    '模块采用推拉结合的 Feed 分发策略，'
    '通过 UserInbox 收件箱模式和大V 配置动态切换机制，在系统吞吐与数据实时性间取得平衡；'
    '集成 Elasticsearch 实现高效全文搜索，覆盖多业务类型的统一检索能力；'
    '同时埋点采集用户行为数据，为后续推荐算法优化提供数据基础。'
    '模块共提供 7 个 REST 接口，完整支撑社区社交全链路业务。'
)
comm_outlook_title = '展望'
comm_outlook_content = (
    '接入协同过滤推荐算法，基于用户行为数据实现个性化动态推荐排序；\n'
    '拓展话题运营能力，新增话题排行榜、话题订阅、话题专属 Feed 等功能；\n'
    '支持视频上传与播放，丰富内容消费形态；\n'
    '引入敏感词过滤与内容风控机制，保障社区内容安全合规。'
)

conclusion_msg_title = '5.4 消息与私信模块结论与展望'
conclusion_msg_content = (
    '消息与私信模块构建了平台完整的用户通信与系统通知体系，覆盖私信聊天、消息通知、系统公告三大场景。'
    '基于 Spring WebSocket 实现实时双向通信，结合 RocketMQ 异步消息队列实现通知的可靠分发与流量削峰；'
    '公告系统支持定时发布、置顶、过期自动下线等完善的生命周期管理；'
    '消息已读/未读追踪与未读计数实时更新，保障用户体验一致性。'
    '模块共提供 5 个 REST 接口及 WebSocket 长连接服务，支撑平台消息实时触达。'
)
msg_outlook_content = (
    '支持消息撤回、已读回执、正在输入等高级聊天功能，提升私信使用体验；\n'
    '接入群聊能力，满足多用户组群沟通场景；\n'
    '新增消息表情回复、图片/文件发送等富媒体消息类型；\n'
    '优化未读消息合并推送策略，减少非核心消息对用户的打扰。'
)

elements_5 = [
    make_paragraph(),
    make_bold_paragraph(conclusion_comm_title),
    make_paragraph(),
    make_bold_paragraph(conclusion_comm_text_title),
    make_paragraph(conclusion_comm_content),
    make_paragraph(),
    make_bold_paragraph(comm_outlook_title),
    make_paragraph(comm_outlook_content),
    make_paragraph(),
    make_bold_paragraph(conclusion_msg_title),
    make_paragraph(),
    make_bold_paragraph(conclusion_comm_text_title),
    make_paragraph(conclusion_msg_content),
    make_paragraph(),
    make_bold_paragraph(comm_outlook_title),
    make_paragraph(msg_outlook_content),
]
insert_elements_after(para_257, elements_5)
print('5.3 & 5.4 inserted!')


# ============================================================
# 4.3 & 4.4
# ============================================================

impl_comm_title = '4.3 社区互动模块实现'
impl_comm_1 = (
    '社区互动模块是系统中' + LQ + '社区动态发布与社交互动' + RQ + '功能的核心业务包，'
    '位于项目 modules/daily 目录下，采用经典的四层分层架构设计。'
    '核心实体 DailyPost 映射 daily_post 数据表，包含用户 ID、内容、图片列表（JSON 数组）、'
    '视频地址、关联话题、地理位置、浏览量/点赞数/评论数/分享数多维统计数据及状态审核位等完整字段；'
    'Topic 映射 topic 表，以热度分 hotScore 支撑话题排序推荐；'
    'Comment 与 LikeRecord 分别映射 comment 和 like_record 表，构成通用评论与点赞系统；'
    'DailyUserBehavior 和 DailyUserInterest 作为推荐数据底座，采集并存储用户兴趣画像；'
    'UserInbox 与 BigVConfig 配合实现推拉结合的 Feed 分发机制。'
)
impl_comm_2 = (
    '（1）Controller 层：DailyPostController @RequestMapping("/pet-system/daily") 注册路由前缀，'
    '对外暴露日常帖子发布（publish）、Feed 流获取（feed）、帖子详情（detail）、'
    '点赞/取消点赞（like）等 REST 接口。CommentController 和 LikeController 分别处理评论与点赞请求。'
    '关键防护手段包括：发布接口标注 @RedisRateLimit 限制单个用户 1 小时 5 次的发布频率，'
    '防止恶意刷屏；点赞接口同样限流避免短时间内大量操作。'
)
impl_comm_3 = (
    '（2）Service 层：DailyPostServiceImpl 是业务逻辑编排核心，注入 DailyPostMapper、'
    'UserMapper、TopicMapper、LikeRecordMapper 等数据访问组件，以及 DistributedCache、SearchDataSyncService 等基础设施。'
    '关键业务实现如下：\n'
    '发布动态（publish）：从 UserContext 获取当前用户 ID，构建 DailyPost 实体后调用 mapper.insert 写入数据库，'
    '随后触发 Feed 分发流程——若发布者粉丝数未超过大V 阈值，则遍历粉丝列表将新动态写入各粉丝的 UserInbox 收件箱（推模式）；'
    '超过阈值的发布者则仅写入标记，由粉丝在拉取时主动合并。'
    '最后通过事件发布器发布 DailyPostPublishedEvent，由 SearchDataSyncListener 监听并同步至 ES 搜索引擎。\n'
    'Feed 流获取（feed）：先从 UserInbox 收件箱分页拉取推模式动态，再异步查询大V 发布者的最新动态做拉模式合并，'
    '以 createTime 倒序归并后返回。同时使用 Redis 缓存收件箱热点数据，'
    '减轻数据库查询压力。\n'
    '全局搜索（UnifiedSearchService）：对接 Elasticsearch 7.x Java Client，构建 bool 查询实现多类型联合检索。'
    '搜索时通过 should 条件同时匹配标题、内容字段，并支持按相关度、时间、热度排序。'
    '搜索结果经统一 DTO 转换后返回，覆盖 daily、pet、activity 三种内容类型的混合检索。'
)
impl_comm_4 = (
    '（3）Mapper 层：DailyPostMapper、TopicMapper、CommentMapper、LikeRecordMapper 继承 MyBatis-Plus BaseMapper，'
    '自动获得通用 CRUD 方法。DailyPostMapper 额外定义 incrementViewCount 等原子自增方法。'
)
impl_comm_5 = (
    '（4）Entity 层：DailyPost 使用 @TableName("daily_post") 绑定数据表，'
    'images 字段以 @TableField(typeHandler = JacksonTypeHandler.class) 实现 Java List 与 JSON 互转；'
    'UserInbox 的 postType 区分动态类型（daily/pet/activity），支撑多类型 Feed 统一存储；'
    'BigVConfig 的 pushThreshold 字段配置大V 推拉模式切换阈值。'
    '所有实体均通过 @TableField(fill = FieldFill.INSERT / INSERT_UPDATE) 配合 MyMetaObjectHandler 实现时间自动填充。'
)

impl_msg_title = '4.4 消息与私信模块实现'
impl_msg_1 = (
    '消息与私信模块是系统中' + LQ + '用户通信与系统通知' + RQ + '能力的基础设施包，'
    '位于项目 modules/message 目录下，与 common 包中的 WebSocket 配置协同工作。'
    '核心实体 PrivateConversation 映射 private_conversation 表，记录会话双方元信息及未读计数；'
    'PrivateMessage 映射 private_message 表，存储每条消息的内容、发送状态及独立删除标记；'
    'UserMessage 映射 user_message 表，作为系统级消息推送管道，覆盖评论通知、点赞通知等多种场景；'
    'Notice 与 NoticeReadRecord 分别映射 notice 和 notice_read_record 表，构成公告发布与阅读追踪体系。'
)
impl_msg_2 = (
    '（1）Controller 层：PrivateMessageController @RequestMapping("/pet-system/private-message") 注册路由前缀，'
    '对外暴露会话列表（conversations）、消息历史（messages）、发送消息（send）等 REST 接口。'
    'MessageController 处理系统消息通知的查询与已读标记；NoticeController 提供公告列表与阅读记录功能。'
    '所有接口通过 JWT 拦截器校验用户身份，确保用户仅能操作自己的会话与消息。'
)
impl_msg_3 = (
    '（2）Service 层：PrivateMessageServiceImpl 是业务逻辑编排核心，'
    '基于 WebSocket + RocketMQ 实现消息的实时推送与可靠投递。关键业务实现如下：\n'
    '发送消息（sendMessage）：先校验会话是否存在，不存在则自动创建 PrivateConversation 并初始化双方未读计数。'
    '插入 PrivateMessage 记录后，通过 RocketMQ 发送异步消息到消息通知队列。'
    '消息消费者从队列中取出消息后，通过 WebSocketHandler 查询对方的 WebSocket 会话是否在线，'
    '在线则实时推送消息内容，离线则等待对方上线后拉取未读消息。'
    '这种' + LQ + '同步落库 + 异步推送' + RQ + '的模式保证了消息不丢失的同时实现了实时触达。\n'
    '消息通知：评论通知、点赞通知、关注通知等业务行为通过 MessageService.sendNotification 方法统一处理。'
    '该方法将通知内容写入 UserMessage 表，同时异步发送 RocketMQ 消息触发 WebSocket 推送。'
    '前端收到推送后更新未读计数徽标，实现全站消息小红点实时同步。\n'
    '公告管理：NoticeService 提供公告的增删改查及定时发布能力。'
    '通过定时任务扫描 scheduledPublishTime 到达待发布状态的公告，自动更新状态并推送通知；'
    '公告过期后自动下线，不再对用户展示。'
)
impl_msg_4 = (
    '（3）WebSocket 层：WebSocketHandler 继承 Spring 的 TextWebSocketHandler，'
    '在 afterConnectionEstablished 中解析 JWT Token 获取用户 ID，建立 userId→WebSocketSession 的映射关系；'
    '在 handleTextMessage 中处理心跳检测（ping/pong）和消息发送请求；'
    '在 afterConnectionClosed 中清理会话映射。'
    '采用 ConcurrentHashMap 存储在线会话，支持多实例部署时通过 Redis Pub/Sub 或 RocketMQ 广播实现跨实例推送。'
)
impl_msg_5 = (
    '（4）Mapper 层：PrivateMessageMapper、PrivateConversationMapper、UserMessageMapper、'
    'NoticeMapper、NoticeReadRecordMapper 继承 MyBatis-Plus BaseMapper，'
    '实现基础 CRUD 与自定义统计查询。'
)
impl_msg_6 = (
    '（5）Entity 层：PrivateMessage 的 senderDelete 和 receiverDelete 字段分别标记双方删除状态，'
    '实现' + LQ + '各自删除不影响对方' + RQ + '的业务语义；'
    'UserMessage 的 type 字段区分消息类型（评论/点赞/关注/系统），'
    'relatedId 关联业务实体便于前端跳转；'
    'Notice 的 scheduledPublishTime 和 expireTime 字段分别控制定时发布与自动过期。'
)

# Add image for community module
pkg_comm_img = r'D:\宠物管理系统\PMS\package_community.png'
img_elem_comm = add_img_to_para(pkg_comm_img, 5.5)

elements_4 = [
    make_paragraph(),
    make_bold_paragraph(impl_comm_title),
    make_paragraph(),
    make_paragraph(impl_comm_1),
    make_paragraph(),
    make_paragraph(impl_comm_2),
    make_paragraph(),
    make_paragraph(impl_comm_3),
    make_paragraph(),
    make_paragraph(impl_comm_4),
    make_paragraph(),
    make_paragraph(impl_comm_5),
    make_paragraph(),
    img_elem_comm,
    make_paragraph(),
    make_paragraph(),
    make_bold_paragraph(impl_msg_title),
    make_paragraph(),
    make_paragraph(impl_msg_1),
    make_paragraph(),
    make_paragraph(impl_msg_2),
    make_paragraph(),
    make_paragraph(impl_msg_3),
    make_paragraph(),
    make_paragraph(impl_msg_4),
    make_paragraph(),
    make_paragraph(impl_msg_5),
    make_paragraph(),
    make_paragraph(impl_msg_6),
]
insert_elements_after(para_240, elements_4)
print('4.3 & 4.4 inserted!')


# ============================================================
# 3.4 包图 & 用例图
# ============================================================

pkg_msg_title = '二、包图'  # 二、包图
pkg_msg_content = (
    '该模块是 PMS 系统中“用户通信与系统通知”能力的基础设施包，位于 com.hongjie.pms.modules.message 包下。'  # 该模块是 PMS 系统中"用户通信与系统通知"能力的基础设施包，位于 com.hongjie.pms.modules.message 包下。
    'controller 层包含 MessageController（消息通知）、PrivateMessageController（私信）和 NoticeController（系统公告）；'  # controller 层包含 MessageController（消息通知）、PrivateMessageController（私信）和 NoticeController（系统公告）；
    'service 层的 MessageService、PrivateMessageService 和 NoticeService 执行业务逻辑编排；'  # service 层的 MessageService、PrivateMessageService 和 NoticeService 执行业务逻辑编排；
    'mapper 层包含 UserMessageMapper、PrivateMessageMapper、PrivateConversationMapper、NoticeMapper 和 NoticeReadRecordMapper，'  # mapper 层包含 UserMessageMapper、PrivateMessageMapper、PrivateConversationMapper、NoticeMapper 和 NoticeReadRecordMapper，
    '基于 MyBatis-Plus 实现对五张数据表的持久化访问；'  # 基于 MyBatis-Plus 实现对五张数据表的持久化访问；
    'entity 层的 UserMessage、PrivateMessage、PrivateConversation、Notice 和 NoticeReadRecord 作为 ORM 实体分别映射对应数据表。'  # entity 层的 UserMessage、PrivateMessage、PrivateConversation、Notice 和 NoticeReadRecord 作为 ORM 实体分别映射对应数据表。
    '模块集成了 WebSocket 实现实时消息推送，结合 RocketMQ 实现异步通知分发，'  # 模块集成了 WebSocket 实现实时消息推送，结合 RocketMQ 实现异步通知分发，
    '确保私信与系统通知的及时性与可靠性。'  # 确保私信与系统通知的及时性与可靠性。
)

use_case_msg_title = '三、用例图'  # 三、用例图
use_case_msg_content = (
    '消息与私信模块聚焦用户间通信与系统通知两大场景，共计 5 个用例：\n'  # 消息与私信模块聚焦用户间通信与系统通知两大场景，共计 5 个用例：
    '私信通信类（2 个）：发送/接收私信、查看私信会话，实现一对一实时聊天；\n'  # 私信通信类（2 个）：发送/接收私信、查看私信会话，实现一对一实时聊天；
    '通知公告类（3 个）：查看系统公告、查看消息通知、标记消息已读，覆盖系统消息全生命周期管理。'  # 通知公告类（3 个）：查看系统公告、查看消息通知、标记消息已读，覆盖系统消息全生命周期管理。
)

# Add image for message module
pkg_msg_img = r'D:\宠物管理系统\PMS\package_message.png'
img_elem_msg = add_img_to_para(pkg_msg_img, 5.5)

elements_msg_design = [
    make_paragraph(),
    make_bold_paragraph(pkg_msg_title),
    make_paragraph(pkg_msg_content),
    make_paragraph(),
    img_elem_msg,
    make_paragraph(),
    make_bold_paragraph(use_case_msg_title),
    make_paragraph(use_case_msg_content),
]
insert_elements_after(para_170, elements_msg_design)
print('3.4 包图 & 用例图 inserted!')


# ============================================================
# 3.3 包图 & 用例图
# ============================================================

pkg_comm_title = '二、包图'  # 二、包图
pkg_comm_content = (
    '该模块是 PMS 系统中“社区动态与内容推送”功能的核心业务包，位于 com.hongjie.pms.modules.daily 及相关包中。'  # 该模块是 PMS 系统中"社区动态与内容推送"功能的核心业务包，位于 com.hongjie.pms.modules.daily 及相关包中。
    'controller 层包含 DailyPostController（日常帖子发布与管理）、CommentController（评论管理）和 LikeController（点赞操作）；'  # controller 层包含 DailyPostController（日常帖子发布与管理）、CommentController（评论管理）和 LikeController（点赞操作）；
    'service 层的 DailyPostService、DailyRecommendService、CommentService、LikeService 和 UnifiedSearchService 协同完成业务逻辑编排、个性化推荐及全局搜索等功能；'  # service 层的 DailyPostService、DailyRecommendService、CommentService、LikeService 和 UnifiedSearchService 协同完成业务逻辑编排、个性化推荐及全局搜索等功能；
    'mapper 层的 DailyPostMapper、TopicMapper、CommentMapper、LikeRecordMapper 基于 MyBatis-Plus 提供数据持久化访问；'  # mapper 层的 DailyPostMapper、TopicMapper、CommentMapper、LikeRecordMapper 基于 MyBatis-Plus 提供数据持久化访问；
    'entity 层的 DailyPost、Topic、DailyTopicRel、Comment、LikeRecord、DailyUserBehavior、DailyUserInterest、UserInbox 和 BigVConfig 作为 ORM 实体分别映射对应数据表。'  # entity 层的 DailyPost、Topic、DailyTopicRel、Comment、LikeRecord、DailyUserBehavior、DailyUserInterest、UserInbox 和 BigVConfig 作为 ORM 实体分别映射对应数据表。
    '模块采用推拉结合的 Feed 分发模式，通过 UserInbox 实现个性化推荐，支持大V 写扩散优化，构成完整的社区互动闭环。'  # 模块采用推拉结合的 Feed 分发模式，通过 UserInbox 实现个性化推荐，支持大V 写扩散优化，构成完整的社区互动闭环。
)

use_case_comm_title = '三、用例图'  # 三、用例图
use_case_comm_content = (
    '社区互动模块围绕内容消费与社交反馈两条主线设计，共计 7 个用例：\n'  # 社区互动模块围绕内容消费与社交反馈两条主线设计，共计 7 个用例：
    '内容消费类（3 个）：发布每日动态、浏览动态流、浏览首页 Feed，构成用户从生产到消费的内容闭环；\n'  # 内容消费类（3 个）：发布每日动态、浏览动态流、浏览首页 Feed，构成用户从生产到消费的内容闭环；
    '社交反馈类（3 个）：点赞/取消点赞、发表评论、删除评论，支撑用户间的轻量化互动；\n'  # 社交反馈类（3 个）：点赞/取消点赞、发表评论、删除评论，支撑用户间的轻量化互动；
    '话题互动类（1 个）：话题互动，提供热门话题浏览与搜索功能。'  # 话题互动类（1 个）：话题互动，提供热门话题浏览与搜索功能。
)

# Add image for community design
pkg_comm_design_img = r'D:\宠物管理系统\PMS\package_community.png'
img_elem_comm_design = add_img_to_para(pkg_comm_design_img, 5.5)

elements_comm_design = [
    make_paragraph(),
    make_bold_paragraph(pkg_comm_title),
    make_paragraph(pkg_comm_content),
    make_paragraph(),
    img_elem_comm_design,
    make_paragraph(),
    make_bold_paragraph(use_case_comm_title),
    make_paragraph(use_case_comm_content),
]
insert_elements_after(para_152, elements_comm_design)
print('3.3 包图 & 用例图 inserted!')


# ============================================================
# Section 2 requirements
# ============================================================

req_comm_title = '社区互动模块'
req_comm_intro = (
    '社区互动是平台社交属性的核心载体，承载用户内容生产、社交互动与信息分发的完整闭环，'
    '实现从动态发布、Feed 推荐、点赞评论到话题运营的全链路社交体验，具体业务需求包含：'
)
req_comm_items = (
    '（1）. 每日动态发布：用户可发布图文或视频动态内容，支持自定义地理位置标记、'
    '关联热门话题，丰富内容表达形式。\n'
    '（2）. 个性化 Feed 推荐：基于用户关注关系和行为数据，采用推拉结合的分发策略，'
    '为用户生成个性化首页信息流，兼顾实时性与系统吞吐。\n'
    '（3）. 点赞互动：支持用户对动态、评论等内容一键点赞或取消点赞，'
    '实时更新点赞计数，增强社交反馈即时性。\n'
    '（4）. 评论互动：支持用户对动态发表评论，以及针对已有评论进行嵌套回复，'
    '构建有深度的社区讨论氛围。\n'
    '（5）. 话题运营：支持动态关联话题标签，提供热门话题浏览与搜索功能，'
    '引导社区内容聚合与传播。\n'
    '（6）. 全局搜索：基于 Elasticsearch 实现跨内容类型的全文检索，'
    '覆盖动态、帖子、活动等多业务类型的统一搜索体验。\n'
    '（7）. 用户行为采集：埋点采集用户在社区中的浏览、点赞、评论、分享等行为数据，'
    '为推荐算法优化和平台运营分析提供数据支撑。'
)

req_msg_title = '消息与私信模块'
req_msg_intro = (
    '消息与私信是平台用户沟通与系统触达的核心基础能力，'
    '覆盖用户间实时私信聊天、系统消息通知与平台公告管理三大场景，'
    '实现信息的高可靠、低延迟传递，具体业务需求包含：'
)
req_msg_items = (
    '（1）. 实时私信聊天：支持用户之间进行一对一实时文字聊天，'
    '消息发送后对方可立即接收，保障沟通实时性。\n'
    '（2）. 私信会话管理：支持查看历史会话列表及聊天记录，'
    '会话按最后消息时间倒序排列，支持双方各自独立删除消息且不影响对方。\n'
    '（3）. 未读消息提醒：实时统计各会话的未读消息数，'
    '在会话列表和导航栏同步展示未读计数，帮助用户快速定位新消息。\n'
    '（4）. 系统消息通知：集成点赞、评论、关注、报名等业务行为的消息通知能力，'
    '用户可查看通知详情并标记已读。\n'
    '（5）. 平台公告管理：管理员可发布系统级公告，支持置顶展示、定时发布与自动过期；'
    '用户可查看公告列表，系统追踪已读/未读状态。\n'
    '（6）. 高可靠消息推送：基于 WebSocket 长连接实现实时推送，'
    '结合 RocketMQ 异步消息队列保证推送可靠性，'
    '离线上线后自动拉取离线期间遗漏的消息。'
)

elements_2 = [
    make_paragraph(),
    make_bold_paragraph(req_comm_title),
    make_paragraph(),
    make_paragraph(req_comm_intro),
    make_paragraph(req_comm_items),
    make_paragraph(),
    make_bold_paragraph(req_msg_title),
    make_paragraph(),
    make_paragraph(req_msg_intro),
    make_paragraph(req_msg_items),
]
insert_elements_after(para_73, elements_2)
print('Section 2 requirements inserted!')

doc.save(DST)
print('Document saved successfully!')
print(f'Output: {DST}')